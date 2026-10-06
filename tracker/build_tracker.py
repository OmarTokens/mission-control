"""Mission Control tracker builder.

Builds a Google-Sheets-ready workbook (.xlsx) from your holdings + targets CSVs:
  Dashboard · Holdings · Control Center · Value Sleeve · Watchlist · Daily Log · History

Prices are live once imported into Google Sheets (GOOGLEFINANCE formulas). The CSV price is only a fallback.

  python3 tracker/build_tracker.py                       # demo data -> out/Mission-Control-Tracker.xlsx
  python3 tracker/build_tracker.py --dir my_data --out out/my-tracker.xlsx

Import: Google Drive -> New -> File upload -> open with Google Sheets.
Then paste tracker/apps_script/MissionControl.gs into Extensions -> Apps Script and run mcSetup() once (adds charts).
"""
import argparse, csv, os, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import CellIsRule, FormulaRule

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ap = argparse.ArgumentParser()
ap.add_argument('--dir', default=os.path.join(ROOT, 'demo'), help='folder with holdings.csv, targets.csv, value_sleeve.csv, history.csv')
ap.add_argument('--out', default=os.path.join(ROOT, 'out', 'Mission-Control-Tracker.xlsx'))
ap.add_argument('--agent', default='Agent', help='what you call your research agent (shown in headings)')
a = ap.parse_args()
rd = lambda f: list(csv.DictReader(open(os.path.join(a.dir, f)))) if os.path.exists(os.path.join(a.dir, f)) else []
HOLD, TARG, VALS, HIST = rd('holdings.csv'), rd('targets.csv'), rd('value_sleeve.csv'), rd('history.csv')
AG = a.agent

INK, MUTED, LINE = "1F2937", "6B7280", "E5E7EB"
fill = lambda c: PatternFill("solid", fgColor=c)
F_DARK, F_SUB, F_TILE, F_GRN, F_RED, F_YEL, F_EDIT = map(fill, ["111827", "F3F4F6", "EEF2FF", "DCFCE7", "FEE2E2", "FEF3C7", "FFF7D6"])
B = Border(bottom=Side(style="thin", color=LINE))
WR = Alignment(wrap_text=True, vertical="top")
USD, USD2, PCT, PCT0 = '$#,##0', '$#,##0.00', '0.0%', '0%'
wb = Workbook()

def put(ws, r, c, v, fmt=None, bold=False, fl=None, color=INK, size=10, wrap=False, border=False):
    x = ws.cell(r, c, v); x.font = Font(name="Roboto", bold=bold, color=color, size=size)
    if fmt: x.number_format = fmt
    if fl: x.fill = fl
    if wrap: x.alignment = WR
    if border: x.border = B
    return x
def bar(ws, r, c1, c2, text):
    for c in range(c1, c2 + 1): ws.cell(r, c).fill = F_DARK
    put(ws, r, c1, text, bold=True, color="FFFFFF", fl=F_DARK, size=11)
def header(ws, r, c1, labels):
    for i, h in enumerate(labels): put(ws, r, c1 + i, h, bold=True, fl=F_SUB, border=True, wrap=True)
def widths(ws, w):
    for i, x in enumerate(w, 1): ws.column_dimensions[L(i)].width = x
def fmtcells(ws, r, cols, fmts=None, bold=()):
    for i, c in enumerate(cols):
        ws.cell(r, c).border = B; ws.cell(r, c).font = Font(name="Roboto", size=10, bold=c in bold)
        if fmts and fmts[i]: ws.cell(r, c).number_format = fmts[i]

# ============================================================ CONTROL CENTER
cc = wb.active; cc.title = "Control Center"
widths(cc, [30, 14, 44, 44])
put(cc, 1, 1, f"CONTROL CENTER: the dials your {AG.lower()} works from", bold=True, size=16)
put(cc, 2, 1, "Yellow cells are yours to change (or ask your agent to change them; it should show the impact first). Everything on the Dashboard is measured against this page.", color=MUTED, size=9)
bar(cc, 4, 1, 4, "1 · TARGET MIX (household)")
header(cc, 5, 1, ["Sleeve", "Target %", "What's in it", "If you raise it…"])
RAISE = {"US": "More growth, deeper drops", "Intl": "More diversification away from US mega-caps", "Gold": "Better crash cushion, low long-run return",
         "T-bills": "Safer, lower return (~T-bill yield)", "Satellites": "More upside and much bigger swings", "Cash": "Cash drag"}
for i, t in enumerate(TARG):
    r = 6 + i
    put(cc, r, 1, t['sleeve'], bold=True, border=True); put(cc, r, 2, float(t['target']), PCT0, fl=F_EDIT, border=True)
    put(cc, r, 3, t.get('what_it_holds', ''), wrap=True, border=True); put(cc, r, 4, RAISE.get(t['sleeve'], ''), wrap=True, border=True)
TN = 6 + len(TARG) - 1
put(cc, TN + 1, 1, "Total (must be 100%)", bold=True); put(cc, TN + 1, 2, f"=SUM(B6:B{TN})", PCT0, bold=True)
TGT = {t['sleeve']: f"'Control Center'!$B${6+i}" for i, t in enumerate(TARG)}

D0 = TN + 4
bar(cc, D0 - 1, 1, 4, "2 · RULES & DIALS")
header(cc, D0, 1, ["Dial", "Current", "What it controls", "Notes"])
DIALS = [("Drift band (pts)", 0.05, "If a sleeve is this far off target, the daily note offers 2 options: sell vs fund from new money", ""),
         ("Value zone: min drop from 52w high", 0.30, "Watchlist/Value Sleeve flag VALUE ZONE when at least this far below the 52-week high…", ""),
         ("Value zone: max P/E", 20, "…AND the P/E is at or below this (ETFs flag on price only)", ""),
         ("Watch: min drop from 52w high", 0.20, "WATCH flag threshold", ""),
         ("Loss-harvest trigger", -0.20, "Value Sleeve flags HARVEST? when a name is down this much", "Mind the 30-day wash-sale rule across ALL accounts"),
         ("Satellite cap per name", 5000, "Max $ in any one satellite", ""),
         ("Exclusions", "e.g. sectors you never want", "Agent never recommends these", ""),
         ("Email recipients", "you@example.com", "Daily + weekly notes go only here", "")]
DIAL = {}
for i, (d, v, what, note) in enumerate(DIALS):
    r = D0 + 1 + i; DIAL[d] = f"'Control Center'!$B${r}"
    f = USD if isinstance(v, int) and v > 100 else (PCT0 if isinstance(v, float) else None)
    put(cc, r, 1, d, bold=True, border=True, wrap=True); put(cc, r, 2, v, f, fl=F_EDIT, border=True, wrap=True)
    put(cc, r, 3, what, wrap=True, border=True); put(cc, r, 4, note, wrap=True, border=True)
S0 = D0 + len(DIALS) + 3
bar(cc, S0, 1, 4, "3 · GLOSSARY")
for i, (k, v) in enumerate([("Sleeve", "A bucket of the portfolio with one job (US stocks, international, gold, T-bills, satellites, cash)."),
                            ("Drift", "How far a sleeve has moved away from its target % because prices moved."),
                            ("P/E", "Price ÷ yearly earnings per share. Lower = cheaper (compare to the stock's own history)."),
                            ("52w high", "Highest price in the last year. 'Off high' = how far below that peak."),
                            ("Tax-loss harvest", "Selling a loser to book a loss that offsets gains you already took. No rebuy of the same thing within 30 days in any account."),
                            ("Satellite", "A small, high-upside bet outside the core.")]):
    r = S0 + 1 + i; put(cc, r, 1, k, bold=True, border=True); put(cc, r, 2, v, wrap=True, border=True)
    cc.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4); cc.row_dimensions[r].height = 28
cc.freeze_panes = "A4"

# ============================================================ HOLDINGS
ho = wb.create_sheet("Holdings")
widths(ho, [26, 11, 12, 12, 11, 11, 11, 13, 12, 9, 40])
put(ho, 1, 1, "HOLDINGS: every position in every account", bold=True, size=16)
put(ho, 2, 1, "Yellow cells (shares, avg cost) are the only hand-entered numbers; update them from a broker screenshot. CASH / CRYPTO / sandbox rows: shares = dollar value. Prices are live.", color=MUTED, size=9)
header(ho, 4, 1, ["Account", "Type", "Ticker", "Sleeve", "Shares", "Avg cost", "Price", "Value", "Gain $", "Gain %", "Note"])
H0 = 5; HN = H0 + len(HOLD) - 1 + 15
for i in range(HN - H0 + 1):
    r = H0 + i; h = HOLD[i] if i < len(HOLD) else None
    if h:
        tk = h['ticker']; fixed = tk in ('CASH', 'CRYPTO') or '-' in tk
        put(ho, r, 1, h['account'], border=True); put(ho, r, 2, h['type'], border=True); put(ho, r, 3, tk, bold=True, border=True); put(ho, r, 4, h['sleeve'], border=True)
        put(ho, r, 5, float(h['shares']), USD if fixed else '#,##0.00', fl=F_EDIT, border=True)
        put(ho, r, 6, float(h['avg_cost']) if h['avg_cost'] else None, USD2, fl=F_EDIT, border=True)
        put(ho, r, 11, h.get('note', ''), color=MUTED, size=9, border=True, wrap=True)
        fb = float(h['price'] or 1)
    else:
        for c in range(1, 7): put(ho, r, c, None, fl=F_EDIT if c in (5, 6) else None, border=True)
        fb = 0
    ho.cell(r, 7).value = (f'=IF($C{r}="",,IF(OR($C{r}="CASH",$C{r}="CRYPTO",ISNUMBER(SEARCH("-",$C{r}))),1,'
                           f'IFERROR(GOOGLEFINANCE($C{r},"price"),{fb})))')
    ho.cell(r, 8).value = f'=IF($C{r}="",,ROUND(N($E{r})*N($G{r}),2))'
    ho.cell(r, 9).value = f'=IF(N($F{r})=0,,ROUND($H{r}-$E{r}*$F{r},2))'
    ho.cell(r, 10).value = f'=IF(N($F{r})=0,,$G{r}/$F{r}-1)'
    fmtcells(ho, r, (7, 8, 9, 10), (USD2, USD, USD, PCT))
put(ho, HN + 2, 7, "TOTAL", bold=True); put(ho, HN + 2, 8, f"=SUM(H{H0}:H{HN})", USD, bold=True); put(ho, HN + 2, 9, f"=SUM(I{H0}:I{HN})", USD, bold=True)
for op, col in (("lessThan", "B91C1C"), ("greaterThan", "15803D")):
    ho.conditional_formatting.add(f"I{H0}:J{HN}", CellIsRule(operator=op, formula=["0"], font=Font(color=col)))
put(ho, HN + 4, 1, "Sleeve must match a row on Control Center (" + " · ".join(t['sleeve'] for t in TARG) + "). Type is Taxable or Retirement.", color=MUTED, size=9)
ho.freeze_panes = "A5"
rng = lambda c: f"Holdings!${c}${H0}:${c}${HN}"
HR, HT, HTK, HS, HV, HG = rng('A'), rng('B'), rng('C'), rng('D'), rng('H'), rng('I')

# ============================================================ VALUE SLEEVE
vs = wb.create_sheet("Value Sleeve")
widths(vs, [8, 22, 11, 11, 11, 10, 13, 12, 9, 10, 12, 9, 40, 34])
put(vs, 1, 1, "VALUE SLEEVE: beaten-down quality names you review monthly", bold=True, size=16)
put(vs, 2, 1, "Example rows only, not recommendations. Zone = at least the Control Center drop below the 52-week high. Owned shares are pulled from Holdings automatically.", color=MUTED, size=9)
header(vs, 4, 1, ["Ticker", "Name", "Price", "52w high", "Off high", "P/E", "Status", "Planned $", "Shares", "Avg cost", "Value", "Gain %", "Why it's cheap / thesis", "Main risk"])
V0 = 5; VN = V0 + max(len(VALS), 10) - 1
for i in range(VN - V0 + 1):
    r = V0 + i; v = VALS[i] if i < len(VALS) else {}
    put(vs, r, 1, v.get('ticker'), bold=True, border=True, fl=F_EDIT); put(vs, r, 2, v.get('name'), border=True)
    put(vs, r, 8, float(v['planned_usd']) if v.get('planned_usd') else None, USD, fl=F_EDIT, border=True)
    put(vs, r, 13, v.get('thesis'), wrap=True, border=True, size=9); put(vs, r, 14, v.get('risk'), wrap=True, border=True, size=9)
    vs.cell(r, 3).value = f'=IF($A{r}="",,IFERROR(GOOGLEFINANCE($A{r},"price"),))'
    vs.cell(r, 4).value = f'=IF($A{r}="",,IFERROR(GOOGLEFINANCE($A{r},"high52"),))'
    vs.cell(r, 5).value = f'=IF(N($D{r})=0,,$C{r}/$D{r}-1)'
    vs.cell(r, 6).value = f'=IF($A{r}="",,IFERROR(GOOGLEFINANCE($A{r},"pe"),))'
    vs.cell(r, 7).value = (f'=IF(N($C{r})=0,,IF(AND(N($L{r})<>0,$L{r}<={DIAL["Loss-harvest trigger"]}),"HARVEST?",'
                           f'IF($E{r}<=-{DIAL["Value zone: min drop from 52w high"]},"IN ZONE",IF($E{r}<=-{DIAL["Watch: min drop from 52w high"]},"NEAR","ABOVE"))))')
    vs.cell(r, 9).value = f'=IF($A{r}="",,SUMIF({HTK},$A{r},{rng("E")}))'
    vs.cell(r, 10).value = f'=IFERROR(SUMPRODUCT(({HTK}=$A{r})*{rng("E")}*{rng("F")})/$I{r},)'
    vs.cell(r, 11).value = f'=IF(N($I{r})=0,,ROUND($I{r}*$C{r},2))'
    vs.cell(r, 12).value = f'=IF(N($J{r})=0,,$C{r}/$J{r}-1)'
    fmtcells(vs, r, (3, 4, 5, 6, 7, 9, 10, 11, 12), (USD2, USD2, PCT, '0.0', None, '#,##0.##', USD2, USD, PCT), bold=(7,))
    vs.row_dimensions[r].height = 30
put(vs, VN + 1, 7, "TOTAL", bold=True); put(vs, VN + 1, 8, f"=SUM(H{V0}:H{VN})", USD, bold=True); put(vs, VN + 1, 11, f"=SUM(K{V0}:K{VN})", USD, bold=True)
G = f"G{V0}:G{VN}"
for val, fl_ in (("IN ZONE", F_GRN), ("NEAR", F_YEL), ("HARVEST?", F_RED)):
    vs.conditional_formatting.add(G, FormulaRule(formula=[f'G{V0}="{val}"'], fill=fl_))
vs.freeze_panes = "C5"

# ============================================================ WATCHLIST
wl = wb.create_sheet("Watchlist")
widths(wl, [8, 28, 18, 10, 10, 10, 10, 9, 9, 10, 13, 30])
put(wl, 1, 1, "WATCHLIST: names you don't need to own, scanned for value", bold=True, size=16)
put(wl, 2, 1, "Live prices. VALUE ZONE / WATCH thresholds come from Control Center. Add or remove tickers freely (column A); copy a row's formulas down for new rows.", color=MUTED, size=9)
header(wl, 4, 1, ["Ticker", "Name", "Group", "Price", "52w low", "52w high", "Off high", "Range pos.", "P/E", "Mkt cap $B", "Flag", "Agent note"])
U = {"Semis & hardware": "NVDA AVGO AMD MU TSM ASML AMAT LRCX KLAC QCOM TXN INTC DELL ANET CSCO",
     "Software & internet": "MSFT GOOGL META AMZN AAPL ORCL CRM NOW ADBE INTU PANW CRWD SHOP UBER BKNG SPOT",
     "Media & telecom": "NFLX DIS CMCSA VZ T EA",
     "Healthcare": "LLY NVO UNH ELV CI CVS ABT MDT BSX SYK ISRG TMO DHR JNJ PFE MRK BMY ABBV AMGN GILD REGN VRTX",
     "Staples": "PEP KO PG CL KMB HSY MDLZ GIS STZ KR COST WMT TGT",
     "Discretionary": "NKE LULU DECK SBUX MCD CMG HD LOW TJX TSLA MAR",
     "Industrials": "UPS FDX UNP CAT DE ETN EMR WM URI VRT",
     "Energy & materials": "XOM CVX COP EOG SLB LIN SHW DOW LYB NUE FCX NEM",
     "Utilities & REITs": "NEE DUK SO CEG VST PLD AMT EQIX O",
     "Payments & exchanges": "V MA PYPL SPGI MCO ICE CME",
     "ETFs": "SPY QQQ VTI VXUS SMH XLV XLE XLU IJR GLDM SGOV SCHD AVUV VWO INDA EWJ"}
r = 5
for grp, s in U.items():
    for t in s.split():
        put(wl, r, 1, t, bold=True, border=True); put(wl, r, 3, grp, size=9, border=True, color=MUTED)
        for c, attr in ((2, "name"), (4, "price"), (5, "low52"), (6, "high52"), (9, "pe")):
            wl.cell(r, c).value = f'=IFERROR(GOOGLEFINANCE($A{r},"{attr}"),)'
        wl.cell(r, 7).value = f'=IF(N($F{r})=0,,$D{r}/$F{r}-1)'
        wl.cell(r, 8).value = f'=IF(N($F{r})-N($E{r})<=0,,($D{r}-$E{r})/($F{r}-$E{r}))'
        wl.cell(r, 10).value = f'=IFERROR(GOOGLEFINANCE($A{r},"marketcap")/1E9,)'
        pe_ok = "TRUE" if grp == "ETFs" else f'AND(N($I{r})>0,$I{r}<={DIAL["Value zone: max P/E"]})'
        wl.cell(r, 11).value = (f'=IF(N($D{r})=0,,IF(AND($G{r}<=-{DIAL["Value zone: min drop from 52w high"]},{pe_ok}),"VALUE ZONE",'
                                f'IF($G{r}<=-{DIAL["Watch: min drop from 52w high"]},"WATCH","—")))')
        fmtcells(wl, r, (2, 4, 5, 6, 7, 8, 9, 10, 11, 12), (None, USD2, USD2, USD2, PCT, PCT0, '0.0', '#,##0', None, None), bold=(11,))
        r += 1
WLN = r - 1
wl.conditional_formatting.add(f"K5:K{WLN}", FormulaRule(formula=['K5="VALUE ZONE"'], fill=F_GRN))
wl.conditional_formatting.add(f"K5:K{WLN}", FormulaRule(formula=['K5="WATCH"'], fill=F_YEL))
wl.freeze_panes = "C5"; wl.auto_filter.ref = f"A4:L{WLN}"

# ============================================================ DAILY LOG
dl = wb.create_sheet("Daily Log")
widths(dl, [12, 8, 46, 46, 34, 34, 34])
put(dl, 1, 1, "DAILY LOG: every scan leaves a note here (newest on top)", bold=True, size=16)
put(dl, 2, 1, "Written by your agent's scans (prompts/daily-scan.md). Research only: you decide every trade.", color=MUTED, size=9)
header(dl, 4, 1, ["Date", "Scan", "Headline", "Ideas (buy / hold / trim)", "Buy-zone alerts", "Rebalance & tax", "Risks"])
for c, v in enumerate([dt.date.today(), "Setup", "Demo tracker created. Replace demo holdings with yours, then run the first scan.",
                       "None yet.", "See Value Sleeve and Watchlist flags.", "Compare Dashboard gaps with the drift band.", "None logged."], 1):
    put(dl, 5, c, v, 'yyyy-mm-dd' if c == 1 else None, wrap=True, border=True, size=9 if c > 2 else 10)
dl.row_dimensions[5].height = 60; dl.freeze_panes = "A5"

# ============================================================ HISTORY
hi = wb.create_sheet("History")
widths(hi, [12, 14, 12, 12, 12, 30])
put(hi, 1, 1, "HISTORY: one row per after-close scan (feeds the timeline chart)", bold=True, size=14)
header(hi, 3, 1, ["Date", "Total", "Benchmark", "You (index)", "Benchmark (index)", "Note"])
for i, h in enumerate(HIST):
    r = 4 + i
    put(hi, r, 1, dt.date.fromisoformat(h['date']), 'yyyy-mm-dd'); put(hi, r, 2, float(h['total']), USD)
    put(hi, r, 3, float(h['spy']), '0.00'); put(hi, r, 6, h.get('note', ''), size=9, color=MUTED)
for r in range(4, 400):
    hi.cell(r, 4).value = f'=IF($B{r}="",NA(),$B{r}/$B$4*100)'; hi.cell(r, 4).number_format = '0.0'
    hi.cell(r, 5).value = f'=IF($C{r}="",NA(),$C{r}/$C$4*100)'; hi.cell(r, 5).number_format = '0.0'

# ============================================================ DASHBOARD
db = wb.create_sheet("Dashboard", 0)
widths(db, [22, 13, 13, 13, 14, 13, 3, 16, 13, 13, 13, 13, 3, 14, 12, 12, 12, 12])
put(db, 1, 1, "MISSION CONTROL: household dashboard", bold=True, size=20)
db.cell(2, 1).value = '="Live prices · last scan: "&TEXT(INDIRECT("\'Daily Log\'!A5"),"ddd d mmm")&" ("&INDIRECT("\'Daily Log\'!B5")&")"'
db.cell(2, 1).font = Font(name="Roboto", size=9, color=MUTED)
dry = " + ".join(f'SUMIF({HS},"{s}",{HV})' for s in ("Cash", "T-bills"))
TILES = [("TOTAL VALUE", f"=SUM({HV})", '="household, live"'),
         ("TAXABLE", f'=SUMIF({HT},"Taxable",{HV})', f'=TEXT(SUMIF({HT},"Taxable",{HV})/SUM({HV}),"0%")&" of total"'),
         ("RETIREMENT", f'=SUMIF({HT},"Retirement",{HV})', f'=TEXT(SUMIF({HT},"Retirement",{HV})/SUM({HV}),"0%")&" of total"'),
         ("CASH + T-BILLS", f"={dry}", f'=TEXT(({dry})/SUM({HV}),"0%")&" not in stocks"'),
         ("UNREALIZED GAIN", f"=SUM({HG})", '="on positions with known cost"'),
         ("VALUE SLEEVE", f"=SUM('Value Sleeve'!$K${V0}:$K${VN})", f"=COUNTIF('Value Sleeve'!$G${V0}:$G${VN},\"IN ZONE\")&\" in buy zone\"")]
for (lab, f, sub), c in zip(TILES, [1, 3, 5, 8, 10, 12]):
    for rr in (4, 5, 6):
        for c2 in (c, c + 1): db.cell(rr, c2).fill = F_TILE
    put(db, 4, c, lab, bold=True, color=MUTED, size=8, fl=F_TILE)
    x = db.cell(5, c, f); x.number_format = USD; x.font = Font(name="Roboto", bold=True, size=15, color=INK)
    y = db.cell(6, c, sub); y.font = Font(name="Roboto", size=8, color=MUTED)

bar(db, 8, 1, 12, "LATEST SCAN")
for i, (pre, col) in enumerate([("", "C"), ("▸ ", "D"), ("▸ Buy zones: ", "E"), ("▸ Rebalance/tax: ", "F"), ("▸ Risks: ", "G")]):
    rr = 9 + i
    db.cell(rr, 1).value = f"=\"{pre}\"&INDIRECT(\"'Daily Log'!{col}5\")"
    db.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=12)
    db.cell(rr, 1).alignment = WR; db.cell(rr, 1).font = Font(name="Roboto", size=11 if i == 0 else 9, bold=i == 0)
    db.row_dimensions[rr].height = 20 if i == 0 else 28

SL = [t['sleeve'] for t in TARG]
bar(db, 15, 1, 6, "ALLOCATION vs TARGET (household)")
header(db, 16, 1, ["Sleeve", "Now $", "Now %", "Target %", "Gap $ (+buy / −sell)", "Progress"])
for i, s in enumerate(SL):
    r = 17 + i
    put(db, r, 1, s, bold=True, border=True)
    db.cell(r, 2).value = f'=SUMIF({HS},$A{r},{HV})'
    db.cell(r, 3).value = f'=IFERROR(B{r}/SUM({HV}),0)'
    db.cell(r, 4).value = f'={TGT[s]}'
    db.cell(r, 5).value = f'=ROUND(D{r}*SUM({HV})-B{r},-2)'
    db.cell(r, 6).value = (f'=SPARKLINE(C{r},{{"charttype","bar";"max",MAX(D{r},C{r},0.0001);'
                           f'"color1",IF(ABS(C{r}-D{r})<={DIAL["Drift band (pts)"]},"#16a34a","#f59e0b")}})')
    fmtcells(db, r, (2, 3, 4, 5, 6), (USD, PCT, PCT0, '+$#,##0;-$#,##0;—', None))
SLN = 17 + len(SL) - 1
for op, col in (("greaterThan", "15803D"), ("lessThan", "B91C1C")):
    db.conditional_formatting.add(f"E17:E{SLN}", CellIsRule(operator=op, formula=["0"], font=Font(color=col)))

bar(db, 15, 8, 12, "TAXABLE vs RETIREMENT")
header(db, 16, 8, ["Sleeve", "Taxable $", "Taxable %", "Retirement $", "Retirement %"])
for i, s in enumerate(SL):
    r = 17 + i
    put(db, r, 8, s, bold=True, border=True)
    db.cell(r, 9).value = f'=SUMIFS({HV},{HS},$H{r},{HT},"Taxable")'
    db.cell(r, 10).value = f'=IFERROR(I{r}/SUMIF({HT},"Taxable",{HV}),0)'
    db.cell(r, 11).value = f'=SUMIFS({HV},{HS},$H{r},{HT},"Retirement")'
    db.cell(r, 12).value = f'=IFERROR(K{r}/SUMIF({HT},"Retirement",{HV}),0)'
    fmtcells(db, r, (9, 10, 11, 12), (USD, PCT0, USD, PCT0))

A0 = SLN + 3
ACC = list(dict.fromkeys(h['account'] for h in HOLD))
TYP = {h['account']: h['type'] for h in HOLD}
bar(db, A0, 1, 6, "ACCOUNTS")
header(db, A0 + 1, 1, ["Account", "Type", "Value", "% of total", "Cash", "Gain $"])
for i, acc in enumerate(ACC):
    r = A0 + 2 + i
    put(db, r, 1, acc, border=True); put(db, r, 2, TYP[acc], border=True, size=9, color=MUTED)
    db.cell(r, 3).value = f'=SUMIF({HR},$A{r},{HV})'
    db.cell(r, 4).value = f'=IFERROR(C{r}/SUM({HV}),0)'
    db.cell(r, 5).value = f'=SUMIFS({HV},{HR},$A{r},{HS},"Cash")'
    db.cell(r, 6).value = f'=SUMIF({HR},$A{r},{HG})'
    fmtcells(db, r, (3, 4, 5, 6), (USD, PCT0, USD, USD))

bar(db, A0, 8, 12, "WATCHLIST RIGHT NOW")
for i, (lab, f) in enumerate([("In VALUE ZONE", f"=COUNTIF(Watchlist!$K$5:$K$400,\"VALUE ZONE\")"),
                              ("On WATCH", f"=COUNTIF(Watchlist!$K$5:$K$400,\"WATCH\")"),
                              ("Names tracked", "=COUNTA(Watchlist!$A$5:$A$400)")]):
    put(db, A0 + 1 + i, 8, lab, bold=True); x = db.cell(A0 + 1 + i, 10, f); x.font = Font(name="Roboto", bold=True, size=14)

# chart data block (cols N-R) — Apps Script mcSetup() draws charts from it
put(db, 15, 14, "chart data (don't edit)", size=8, color=MUTED)
for c, h in zip(range(14, 19), ["Sleeve", "Now", "Target", "Taxable", "Retirement"]): db.cell(16, c).value = h
for i, s in enumerate(SL):
    r = 17 + i
    db.cell(r, 14, s); db.cell(r, 15).value = f"=B{r}"; db.cell(r, 16).value = f"=D{r}*SUM({HV})"
    db.cell(r, 17).value = f"=I{r}"; db.cell(r, 18).value = f"=K{r}"
for rr in range(16, SLN + 1):
    for c in range(14, 19): db.cell(rr, c).font = Font(name="Roboto", size=8, color="9CA3AF"); db.cell(rr, c).number_format = USD if c > 14 else "General"
db.freeze_panes = "A4"

for ws in wb.worksheets: ws.sheet_view.showGridLines = False
os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
wb.save(a.out)
print(f"saved {a.out}: {len(HOLD)} holdings, {len(ACC)} accounts, {WLN-4} watchlist names, chart data Dashboard!N16:R{SLN}")
