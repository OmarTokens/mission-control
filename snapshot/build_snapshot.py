"""Mission Control snapshot builder.

Turns a holdings list + target mix into:
  1. snapshot.html        a full, self-contained dashboard page (open it in a browser or publish it as a Claude artifact)
  2. snapshot_email.html  an email-safe block (inline styles, table bars) to paste at the top of a daily email
  3. snapshot.png         optional screenshot (needs Playwright + Chromium)

Two input modes:
  python3 build_snapshot.py --csv ../demo/holdings.csv --targets ../demo/targets.csv --out out/snapshot.html [--png out/snapshot.png]
  python3 build_snapshot.py --xlsx tracker.xlsx --out out/snapshot.html      # Google Sheet exported as .xlsx (Holdings + Control Center tabs, values computed)

Holdings columns: account, type (Taxable|Retirement), ticker, sleeve, shares, avg_cost, price, note
Sleeves: US, Intl, Gold, T-bills, Satellites, Cash. For CASH / value-only rows use shares = dollar value and price = 1.
"""
import argparse, csv, json, os, datetime as dt

SLEEVES = ['US', 'Intl', 'Gold', 'T-bills', 'Satellites', 'Cash']
ap = argparse.ArgumentParser()
ap.add_argument('--csv'); ap.add_argument('--targets'); ap.add_argument('--xlsx')
ap.add_argument('--out', required=True); ap.add_argument('--png'); ap.add_argument('--asof')
ap.add_argument('--title', default='Mission Control'); ap.add_argument('--sat-cap', type=float, help='max $ per satellite (default: 10%% of the household)')
ap.add_argument('--drift', type=float, default=0.05, help='drift band, e.g. 0.05 = 5 points')
ap.add_argument('--template', default=os.path.join(os.path.dirname(os.path.abspath(__file__)), 'template.html'))
a = ap.parse_args()

rows, targets = [], {}
if a.csv:
    for r in csv.DictReader(open(a.csv)):
        v = float(r['shares']) * float(r['price'] or 1)
        rows.append((r['account'], r['type'], r['ticker'], r['sleeve'], float(r['shares']), v))
    for r in csv.DictReader(open(a.targets)):
        targets[r['sleeve']] = float(r['target'])
else:
    import openpyxl
    wb = openpyxl.load_workbook(a.xlsx, data_only=True)
    H, C = wb['Holdings'], wb['Control Center']
    for r in range(5, H.max_row + 1):
        acc, typ, tk, sl, sh, avg, px, val = (H.cell(r, c).value for c in range(1, 9))
        if tk and isinstance(val, (int, float)) and val: rows.append((acc, typ, tk, 'T-bills' if sl == 'SGOV' else sl, float(sh or 0), float(val)))
    for r in range(6, 20):
        k, v = C.cell(r, 1).value, C.cell(r, 2).value
        if not k or str(k).startswith('Total'): break
        targets['T-bills' if k in ('SGOV', 'JAAA') else k] = targets.get('T-bills' if k in ('SGOV', 'JAAA') else k, 0) + float(v or 0)

accts = {}
for acc, typ, tk, sl, sh, v in rows:
    A = accts.setdefault(acc, {'n': acc, 'type': typ, 'src': ('holdings.csv' if a.csv else 'tracker'), 'p': []})
    label = {'CASH': 'Cash', 'CRYPTO': 'Crypto'}.get(tk, tk)
    value_only = tk in ('CASH', 'CRYPTO') or sl == 'Cash' or '-' in tk   # cash, crypto, sandbox rows: shares = dollars
    A['p'].append([label, sl, None if value_only else round(sh, 2), round(v, 2)])
accounts = list(accts.values())
for A in accounts: A['p'].sort(key=lambda p: -p[3])
total = sum(p[3] for A in accounts for p in A['p'])
by = {}
for A in accounts:
    for p in A['p']: by[p[1]] = by.get(p[1], 0) + p[3]

K = lambda v: f"${v/1000:,.1f}K"; P = lambda v: f"{v*100:.1f}%"
items = []
for s in SLEEVES[:-1]:
    now, t = by.get(s, 0) / total, targets.get(s, 0)
    if abs(now - t) > a.drift:
        gap = t * total - by.get(s, 0)
        items.append((abs(gap), ['warn', s, f"{s} is {P(now)} vs a {t*100:.0f}% target, {K(abs(gap))} {'below' if gap > 0 else 'above'} it."]))
items = [i for _, i in sorted(items, key=lambda x: -x[0])]
if by.get('Cash', 0) > 0.01 * total:
    parts = [f"{K(sum(p[3] for p in A['p'] if p[1]=='Cash'))} {A['n']}" for A in accounts if sum(p[3] for p in A['p'] if p[1] == 'Cash') > 0.002 * total]
    items.append(['warn', 'Cash', f"{K(by['Cash'])} uninvested: " + ", ".join(parts) + "."])
cap = a.sat_cap or 0.10 * total
for tk in {p[0] for A in accounts for p in A['p'] if p[1] == 'Satellites'}:
    v = sum(p[3] for A in accounts for p in A['p'] if p[0] == tk)
    if v > cap and tk not in ('Crypto',):
        items.append(['warn', tk, f"{tk} is {K(v)} ({P(v/total)}), above the {K(cap)} per-satellite cap."])

data = {'title': a.title,
        'sub': f"All accounts as of {a.asof or dt.date.today().strftime('%a %d %b %Y')}. Research only; you decide every trade.",
        'targets': targets, 'drift': a.drift, 'accounts': accounts, 'items': items[:6]}
html = open(a.template).read().replace('__DATA__', json.dumps(data))
os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
open(a.out, 'w').write(html)

# ---- email-safe block
COL = {'US': '#1f6f8b', 'Intl': '#4f9d69', 'Gold': '#c9a227', 'T-bills': '#8a96a3', 'Satellites': '#8b5cf6', 'Cash': '#c5cdd5'}
F = lambda v: f"${v:,.0f}"
tt = {}
for A in accounts: tt[A['type']] = tt.get(A['type'], 0) + sum(p[3] for p in A['p'])
dry = by.get('T-bills', 0) + by.get('Cash', 0)
td = 'padding:6px 8px;font:13px Arial,sans-serif;color:#14202b;'
def kpi(l, v, s):
    return (f"<td style='padding:10px 12px;border:1px solid #dde3e8;background:#fff;'><div style='font:11px Arial;color:#5d6b78;text-transform:uppercase'>{l}</div>"
            f"<div style='font:bold 20px Arial;color:#14202b'>{v}</div><div style='font:12px Arial;color:#5d6b78'>{s}</div></td>")
def bar(parts, w=260):
    t = sum(v for _, v in parts) or 1
    return "<table cellpadding='0' cellspacing='0' style='border-collapse:collapse;width:%dpx'><tr>%s</tr></table>" % (w, ''.join(
        f"<td style='background:{c};height:12px;width:{max(1, round(v/t*w))}px;font-size:0;line-height:0'>&nbsp;</td>" for c, v in parts if v > 0))
mx = max(max(by.get(s, 0) / total, targets.get(s, 0)) for s in COL)
mix = ''.join(f"<tr><td style='{td}font-weight:bold'><span style='color:{COL[s]}'>&#9632;</span> {s}</td><td style='{td}'>{bar([(COL[s], by.get(s,0)/total), ('#e8edf1', max(mx - by.get(s,0)/total, 0))])}</td>"
              f"<td style='{td}white-space:nowrap'>{P(by.get(s,0)/total)} vs {targets.get(s,0)*100:.0f}%</td></tr>" for s in COL)
acc = ''
for A in accounts:
    o = {}
    for p in A['p']: o[p[1]] = o.get(p[1], 0) + p[3]
    acc += f"<tr><td style='{td}font-weight:bold'>{A['n']}</td><td style='{td}'>{bar([(COL.get(s,'#999'), o.get(s,0)) for s in COL])}</td><td style='{td}text-align:right'>{F(sum(o.values()))}</td></tr>"
lis = ''.join(f"<li style='margin:0 0 6px;font:13px Arial'><b style='color:#b7791f'>{l}:</b> {t}</li>" for _, l, t in items[:6])
kp = (f"<table cellspacing='6' cellpadding='0' style='width:100%'><tr>{kpi('Household', F(total), f'{len(accounts)} accounts')}"
      f"{kpi('Taxable', F(tt.get('Taxable',0)), P(tt.get('Taxable',0)/total))}</tr><tr>{kpi('Retirement', F(tt.get('Retirement',0)), P(tt.get('Retirement',0)/total))}"
      f"{kpi('Cash + T-bills', F(dry), P(dry/total))}</tr></table>")
email = (f"<div style='max-width:680px;background:#f5f7f8;padding:14px;font-family:Arial,sans-serif'><div style='font:bold 18px Arial;margin-bottom:8px'>{a.title} snapshot</div>"
         f"{kp}<div style='font:bold 12px Arial;color:#5d6b78;margin:12px 0 4px'>MIX VS TARGET</div><table style='background:#fff;border:1px solid #dde3e8;width:100%'>{mix}</table>"
         f"<div style='font:bold 12px Arial;color:#5d6b78;margin:12px 0 4px'>ACCOUNTS</div><table style='background:#fff;border:1px solid #dde3e8;width:100%'>{acc}</table>"
         f"<div style='font:bold 12px Arial;color:#5d6b78;margin:12px 0 4px'>WHAT NEEDS A DECISION</div><ul style='padding-left:18px;margin:0'>{lis}</ul></div>")
open(a.out.replace('.html', '_email.html'), 'w').write(email)
print(json.dumps({'total': round(total), 'by_sleeve': {k: round(v) for k, v in by.items()}}))

if a.png:
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        b = p.chromium.launch(); pg = b.new_page(viewport={'width': 1000, 'height': 900}, device_scale_factor=2)
        pg.set_content("<!doctype html><meta charset='utf-8'><body>" + html, wait_until='networkidle'); pg.wait_for_timeout(600)
        pg.screenshot(path=a.png, full_page=True); b.close()
